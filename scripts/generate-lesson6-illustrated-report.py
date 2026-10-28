from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
SHOTS=ROOT/'docs/evidence/lesson6-real-model/screenshots'
SCRATCH=ROOT/'local_runs/lesson6-report/qa'
SCRATCH.mkdir(parents=True,exist_ok=True)
doc=Document();sec=doc.sections[0]
for style in doc.styles:
 for node in list(style.element.iter()):
  if node.tag == qn('w:pBdr'): node.getparent().remove(node)
for style in doc.styles:
 if style.type == 1:
  style.font.color.rgb=RGBColor(0,0,0)
  style.paragraph_format.keep_with_next=False
  style.paragraph_format.keep_together=False
  style.paragraph_format.widow_control=False
  style.paragraph_format.page_break_before=False
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 s=doc.styles[name];s.font.name='Songti SC';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
 s.font.size=Pt(12 if name=='Normal' else 15 if name=='Heading 1' else 18 if name=='Title' else 12 if name=='Heading 2' else 10.5)
 s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.line_spacing=1.2
 s.paragraph_format.keep_with_next=name in ['Title','Heading 1','Heading 2']
for node in list(sec._sectPr):
 if node.tag == qn('w:docGrid'): sec._sectPr.remove(node)
for name in ['Normal','Caption','Heading 1','Heading 2','Title']:
 st=doc.styles[name]
 st.paragraph_format.line_spacing=Pt(18 if name=='Normal' else 14 if name=='Caption' else 23)
 node=OxmlElement('w:snapToGrid');node.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(node)
doc.styles['Title'].font.bold=True;doc.styles['Heading 1'].font.bold=True
footer=sec.footer.paragraphs[0];footer.alignment=1
r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
def p(t):
 x=doc.add_paragraph(t);x.paragraph_format.first_line_indent=Cm(.84);return x
def image(name,caption,crop=None,maxheight=8.5):
 path=SHOTS/name
 if crop:
  im=Image.open(path);im=im.crop(crop if isinstance(crop,tuple) else (0,0,crop,im.height));path=SCRATCH/name;im.save(path,quality=97)
 im=Image.open(path);w=min(16.4,maxheight*im.width/im.height)
 para=doc.add_paragraph();para.alignment=1;para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.0
 para.add_run().add_picture(str(path),width=Cm(w))
 cap=doc.add_paragraph(caption,style='Caption');cap.alignment=1

def page(title):
 h=doc.add_heading(title,level=1);h.paragraph_format.page_break_before=True
doc.add_paragraph('CampusClaw 第六课项目验证报告',style='Title')
doc.add_paragraph('仓库地址：https://github.com/stanleyIl/forth-lesson')
p('CampusClaw 是一个面向班级教学的 AI 应用平台。教师可以上传课程材料，学生登录后可以检索本班知识库。平台内置知识助手和公式辅导，支持提示词模板、内置与上传的 Skill，以及 MCP 工具接入。')
p('本次使用 glm-5.3-flash 完成验证，接口为 https://llm.goaichat.top/v1。知识助手使用关键词检索，公式辅导通过工具完成计算。报告同时保留网页操作结果和后端实际请求记录。')
p('后端截图通过浏览器打开 HTTP 采集日志获得。日志展开了实际发送的参数、消息和服务返回值，Authorization 密钥不写入采集记录。')
image('01-apps.jpg','图 1 教师登录后的多个 AI 应用')
page('一、提示词模板与技能上传')
p('教师可以在网页中预览模板变量的替换结果，也可以上传 SKILL.md。本次上传“证据摘要”1.0.1 版本，页面提示上传成功，并在注册表中标为 uploaded。随后将它绑定到知识助手。')
image('05-template.jpg','图 2 提示词模板预览与平台管理')
image('07-skill-upload.jpg','图 3 上传技能后的注册表')
page('二、知识库问答与内容溯源')
p('在知识助手中输入“磁铁”，系统选中 1 个本班切片，返回了磁铁性质的回答。网页保留材料名称、切片编号、偏移和原文，可核对回答的依据。')
image('02-knowledge.jpg','图 4 知识库回答及来源引用')
p('后端请求的 system 消息包含模板正文、这一条材料片段，以及两个已绑定技能的指令。请求没有加入其他班级材料或整个知识库。上传技能的指令正文已传入，文件的元数据和存储路径没有传入。')
image('14-backend-knowledge-log.jpg','图 5 实际发送的模板 知识片段和 Skill 指令')
page('三、具体 Skill 的执行过程')
p('在公式辅导中调用要点提取技能，传入关于材料上传、班级检索和来源标注的三句话。网页显示 extracted 3 key points，并展示了三个返回要点。')
image('04-skill-answer.jpg','图 6 要点提取技能的网页执行结果')
p('后端记录中，模型先返回 skill__extract_key_points 和 text 参数。服务端检查绑定关系并校验参数，执行固定动作，再将三个要点作为 role=tool 的消息交回模型。')
image('16-backend-skill-log.jpg','图 7 Skill 调用参数与工具结果回传',crop=(0,290,850,615))
page('四、MCP 工具调用与响应')
p('公式辅导已绑定 calculate 和 lookup_class_schedule。输入“请调用 calculate 计算工具计算 37 × 23”，模型选择计算工具，页面最终显示结果 851。')
image('03-mcp-answer.jpg','图 8 真实模型调用 MCP 后的网页回答')
p('下面的后端响应状态为 200，模型标识为 glm-5.3-flash，finish_reason 为 tool_calls。返回参数为 operation=multiply、a=37、b=23。响应头也保留在同一条记录中。')
image('15-backend-mcp-call-log.jpg','图 9 模型响应头与工具调用参数')
page('五、发送给模型的工具定义')
p('首次请求包含 model、messages、max_tokens=3000、stream=false 和 tool_choice=auto。messages 包含渲染后的 system 提示词及用户问题。tools 只列出当前应用绑定的技能和 MCP 工具。')
image('18-backend-mcp-request.jpg','图 10 首次模型请求和工具列表')
p('每个工具都提供名称、说明和参数结构。计算工具要求 operation、a 和 b，其中 operation 限制为加减乘除或乘方，数值参数也有范围限制。服务端使用同一结构检查模型填入的参数。')
image('19-backend-tool-schema.jpg','图 11 实际传入的计算工具参数结构',crop=(0,150,650,720))
page('六、工具结果如何形成最终回答')
p('MCP 服务执行计算后返回 851。服务端再次请求模型，消息中保留 assistant 的工具调用，并加入 tool 消息。tool_call_id 将计算结果与前面的调用对应起来。')
image('10-backend-tool-return.jpg','图 12 回传模型的计算结果和消息角色',crop=(0,310,850,665))
p('第二次模型响应的 finish_reason 为 stop，回答“根据计算工具结果，37 × 23 = 851”。响应中的 total_tokens 为 933；这只是第二次请求的用量，不是整段会话的总用量。')
image('11-backend-final-response.jpg','图 13 工具执行后的最终模型响应',crop=(0,240,850,710))
page('七、Token 认证与权限检查')
p('不提供登录凭据访问 /api/me，后端返回 401。携带登录取得的 Bearer Token 后返回 200，并给出用户标识、角色和班级。后续请求由服务端解析 Token 来识别用户，不能通过前端传入的 class_id 改变访问范围。')
p('A 班教师访问 B 班材料文件时返回 404。学生尝试上传 Skill 时返回 403。下面是这些接口实际返回的结果，Token 在展示时已隐藏。')
image('17-backend-auth-checks.jpg','图 14 后端认证 班级隔离和学生权限验证',crop=(0,0,900,380),maxheight=12)
p('本次还运行了 npm test、npm run build 和 npx openspec validate --specs。结果为 100 项测试通过，TypeScript 构建通过，14 项 spec 校验通过。')
page('八、五个验证问题的回答')
doc.add_heading('对话参数和提示词模板',level=2)
p('模型请求带有 model、messages、max_tokens 和 stream；有工具的应用还带有 tools 与 tool_choice。提示词模板渲染后作为 system 消息一起传递。模型返回工具调用指令或最终文本。图 5、图 9 至图 13 展示了具体内容。')
doc.add_heading('知识库是否全部传入',level=2)
p('不会。知识助手先检索登录用户所在班级的材料，只把选中的知识片段加入提示词。本次选中 1 个切片，来源是 A班-小学科学课程资料.md。图 4 和图 5 可以相互核对。')
doc.add_heading('Skill 是否全部传入',level=2)
p('需要区分类型。instruction 类型会把已绑定版本的指令正文传入，但不传整个文件的元数据。server_action 类型给模型的是动作名称、说明和参数结构，具体实现留在后端。图 5 显示了上传技能正文，图 10 显示了动作定义。')
doc.add_heading('Skill 如何执行',level=2)
p('模型提出动作调用后，服务端校验应用绑定关系和参数，再执行已注册的动作。结果作为 tool 消息参与后续请求。要点提取动作按句子拆分输入并取限定数量的要点，上传文件不能任意执行代码。图 6 和图 7 展示了这一过程。')
doc.add_heading('MCP 接入后传递什么',level=2)
p('模型收到已绑定工具的名称、说明和参数结构，不直接连接 MCP 服务。模型返回工具名和 JSON 参数，由服务端执行 tools/call，再将结果作为 tool 消息交回模型。图 8 至图 13 对应本次完整调用过程。')
output=ROOT/'output/docx/campusclaw-lesson-6-report.docx';output.parent.mkdir(parents=True,exist_ok=True)
doc.core_properties.title='CampusClaw 第六课项目验证报告';doc.core_properties.author=''
doc.save(output);print(output)
